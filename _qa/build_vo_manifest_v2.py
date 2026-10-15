"""VO_Manifest_v2.xlsx - every VO the game plays now, screen by screen and, inside each question, hint by hint.
Columns: Screen | Screen name | When it plays | Audio ID | Script | Status."""
import json, os, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
html = open(os.path.join(root, "index.html"), encoding="utf-8").read()
card = json.loads(re.search(r'id="cardData">(.*?)</script>', html, re.S).group(1))
A = card["assets"]; AUD = A["audio"]
TEXT = json.load(open(os.path.join(root, "_qa", "review_voices.json"), encoding="utf-8")); TEXT.update(A["audio_text"])

def fid(k): return os.path.splitext(os.path.basename(AUD[k].split("?")[0]))[0]
def status(k):
    f = fid(k)
    if "?v=rec2" in AUD[k]: return "Recorded (voiceovers 2)"
    if f == "vo_review_landing_v2": return "Edited (recording ke tukde jode) - sahi kram mein dobara recording karein"
    if f == "vo_review_story_8a": return "Recorded (edited: अंत का 'भौं, भौं, भौं' हटाया - उसकी जगह SFX)"
    if "?v=rec1" in AUD[k]: return "Recorded"
    return "AI awaaz - recording chahiye"
SCRIPT_FIX = {"vo_review_landing_v2": "नमस्ते दोस्त, मैं हूँ स्विफ़्टी। आज हम एक कहानी पढ़ेंगे—आलू-मालू-कालू।",
              "vo_review_story_8a": "मालू आलू खोज रहा था कि उसे सुनाई दिया,"}

rows = []   # (screen, name, when, id, script, status, kind)
def vo(screen, name, when, k):
    f = fid(k); rows.append([screen, name, when, f, SCRIPT_FIX.get(f, TEXT.get(k, "")), status(k), "vo"])
def note(screen, name, when, ids, script, st="Recorded (upar wale clips dobara)"):
    rows.append([screen, name, when, ids, script, st, "note"])
def head(screen, name):
    rows.append([screen, name, "", "", "", "", "head"])

def tap_question(code, sl, puzzle=False):
    au, d = sl["audio"], sl["data"]; name = sl["prompt_hi"]
    head(code, name)
    q = au.get("question") or au["prompt"]
    vo(code, "", "Sawaal (shuru mein)", q)
    opts = d["options"]
    for n, o in enumerate(opts, 1):
        vo(code, "", f"Option {n}" + (" (sahi)" if o.get("correct") else ""), o["audio"])
    if d.get("hint1_vo"): vo(code, "", "Hint 1 (1st galat)", au["try_again"])
    elif d.get("hint1_question"): note(code, "", "Hint 1 (1st galat)", fid(q), "(sawaal dobara)")
    else: note(code, "", "Hint 1 (1st galat)", "-", "(koi VO nahi - sirf wiggle)", "-")
    ids = " + ".join(([] if d.get("hint2_options_only") else [fid(q)]) + [fid(o["audio"]) for o in opts])
    note(code, "", "Hint 2 (2nd galat)", ids, ("(sirf options, ek-ek karke highlight)" if d.get("hint2_options_only") else "(sawaal, phir options ek-ek karke highlight)"))
    vo(code, "", "Hint 3 (3rd galat - sahi option pulse + haath)", au["reveal"])
    correct = [o for o in opts if o.get("correct")][0]
    note(code, "", "Sahi jawab", fid(correct["audio"]) + " + vo_shabash", "(sahi shabd, phir 'शाबाश!' + confetti)")

def fill_question(code, sl):
    au, d = sl["audio"], sl["data"]
    head(code, sl["prompt_hi"])
    vo(code, "", "Sawaal / heading (shuru mein)", au["prompt"])
    for n, o in enumerate(d["options"], 1): vo(code, "", f"Shabd {n}", o["audio"])
    note(code, "", "Hint 1 (1st galat)", fid(au["prompt"]), "(heading wali line dobara)")
    note(code, "", "Hint 2 (2nd galat)", " + ".join(fid(o["audio"]) for o in d["options"]), "(shabd ek-ek karke highlight)")
    vo(code, "", "Hint 3 (3rd galat - sahi shabd + haath)", au["reveal"])
    note(code, "", "Sahi jawab", "vo_shabash + " + fid(au["correct"]), "('शाबाश!', phir poora vaakya - shabd ek-ek karke highlight)")

# --- title ---
head("Title", "आलू-मालू-कालू (title screen)")
vo("Title", "", "Greeting (screen khulte hi)", "vo_landing")
q_no = 0; p_no = 0; t_no = 1
for sl in card["slides"]:
    t, d, au = sl["type"], sl.get("data", {}), sl.get("audio", {})
    if t == "PHASE_TRANSITION":
        code = {"PT1": "PT1", "PT2": "PT2", "PT3": "PT3"}[sl["id"]]
        cap = d.get("caption", {}).get("text", "")
        head(code, f"Transition - heading: {cap}")
        vo(code, "", f"Swiftie (heading '{cap}' isi line par aati hai)", au["prompt"])
    elif t == "STORY_READALONG":
        t_no += 1; code = f"T{t_no}"
        head(code, f"Story page {d.get('page_no')}")
        if au.get("prompt"): vo(code, "", "Instruction (page khulte hi)", au["prompt"])
        for ln in d["passage"]["lines"]: vo(code, "", "Kahani ki line (mic dabane ke baad)", ln["audio"])
        if d.get("sfx_after"):
            rows.append([code, "", "SFX (line ke turant baad)", d["sfx_after"],
                         "बड़े कुत्ते की असली भौंक — भौं, भौं, भौं (~1.3 सेकंड)। स्रोत: freesound.org/s/784914, CC0",
                         "SFX (asli recording)", "sfx"])
    elif t == "TAP_IN_SCENE":
        q_no += 1; code = f"Q{q_no}"
        head(code, sl["prompt_hi"])
        vo(code, "", "Sawaal (shuru mein)", au["prompt"])
        vo(code, "", "Hint 1 (1st galat)", au["try_again"])
        note(code, "", "Hint 2 (2nd galat)", "vo_o_kaalu + vo_review_identify_correct", "(कालू, मालू - box bari-bari pulse)")
        vo(code, "", "Hint 2 - naam 1", "vo_o_kaalu")
        vo(code, "", "Hint 2 - naam 2", "vo_review_identify_correct")
        vo(code, "", "Hint 3 (3rd galat - sahi box + haath)", au["reveal"])
        note(code, "", "Sahi jawab", fid(au["correct"]) + " + vo_shabash", "('मालू।', phir 'शाबाश!' + confetti)")
    elif t == "EVIDENCE_QA":
        q_no += 1; tap_question(f"Q{q_no}", sl)
    elif t == "PICTURE_PUZZLE":
        head("P", "Picture puzzle")
        vo("P", "", "Puzzle ki instruction", au["prompt"])
        for q in d["questions"]:
            p_no += 1; code = f"P-Q{p_no}"
            (fill_question if q["data"].get("kind") == "fill_blanks" else tap_question)(code, q)
        head("P", "Tukde jodna")
        vo("P", "", "Saare sawaal ke baad", "vo_review_puzzle_assemble")
    elif t == "CELEBRATION":
        head("End", "Celebration")
        vo("End", "", "End screen", au["prompt"])

# shared clip, listed once at the end for the recording team
head("Common", "Har sahi jawab par")
vo("Common", "", "Sahi jawab (confetti ke saath)", "vo_shabash")

# --- write ---
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "VO Manifest"
H = ["Screen", "Screen ka naam", "Kab bajta hai", "Audio ID", "Script", "Status"]
ws.append(H)
hdr = PatternFill("solid", fgColor="1F3864"); sec = PatternFill("solid", fgColor="DCE6F1")
ai = PatternFill("solid", fgColor="FFE699"); ed = PatternFill("solid", fgColor="F8CBAD"); thin = Side(style="thin", color="BFBFBF")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF", size=12); c.fill = hdr; c.alignment = Alignment(horizontal="center", vertical="center")
for r in rows:
    kind = r[6]; ws.append(r[:6]); row = ws[ws.max_row]
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    if kind == "head":
        for c in row: c.fill = sec; c.font = Font(bold=True, color="1F3864")
    elif kind == "note":
        for c in row[2:]: c.font = Font(italic=True, color="595959")
    if r[5].startswith("AI awaaz"): row[5].fill = ai; row[5].font = Font(bold=True)
    if r[5].startswith("Edited"): row[5].fill = ed; row[5].font = Font(bold=True)
for col, w in zip("ABCDEF", (10, 34, 34, 44, 70, 40)): ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
out = os.path.join(root, "VO_Manifest_v2.xlsx"); wb.save(out)
n_vo = sum(1 for r in rows if r[6] == "vo"); n_ai = sum(1 for r in rows if r[5].startswith("AI awaaz"))
print("rows", len(rows), "| VO rows", n_vo, "| AI-voice rows", n_ai)
for r in rows:
    if r[5].startswith(("AI awaaz", "Edited")): print("  ", r[0], r[2], r[3], "|", r[4][:50], "|", r[5][:40])
